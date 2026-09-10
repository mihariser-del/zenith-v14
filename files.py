import os
import io
import re
import uuid
import base64
import tempfile
import subprocess
import httpx
from datetime import datetime, timezone
from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile, File
from fastapi.responses import FileResponse as FastAPIFileResponse
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from database import UploadedFile, get_db
from auth import get_current_user_from_cookie
from limits import check_limit

router = APIRouter(prefix="/api/files", tags=["files"])

UPLOAD_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)

MAX_SIZE = 20 * 1024 * 1024
MAX_TEXT = 50000

ALLOWED_TYPES = {
    "image/png", "image/jpeg", "image/gif", "image/webp", "image/bmp",
    "text/plain", "text/markdown", "text/csv", "text/html", "text/css",
    "text/xml", "text/x-python", "text/x-c", "text/x-c++", "text/javascript",
    "application/json", "application/javascript", "application/xml",
    "application/x-yaml", "application/x-sh", "application/x-httpd-php",
    "application/pdf",
    "application/msword", "application/vnd.ms-excel", "application/vnd.ms-powerpoint",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "audio/mpeg", "audio/wav", "audio/x-wav", "audio/mp4", "audio/m4a",
    "audio/aac", "audio/ogg", "audio/opus", "audio/flac", "audio/x-m4a",
    "audio/x-ms-wma", "audio/webm", "audio/3gpp",
    "video/mp4", "video/webm", "video/quicktime", "video/x-msvideo",
    "video/x-matroska", "video/avi", "video/3gpp", "video/mp2t",
}

EXT_ALLOWLIST = {
    ".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp",
    ".txt", ".md", ".csv", ".json", ".html", ".css", ".xml", ".yaml", ".yml",
    ".py", ".js", ".ts", ".jsx", ".tsx", ".c", ".cpp", ".h", ".java", ".go",
    ".rs", ".php", ".rb", ".sh", ".sql", ".log", ".ini", ".conf", ".toml",
    ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx",
    ".mp3", ".wav", ".m4a", ".aac", ".ogg", ".opus", ".flac", ".wma",
    ".mp4", ".webm", ".mov", ".avi", ".mkv", ".3gp", ".3g2", ".ts",
}


def _file_allowed(content_type, filename) -> bool:
    if not content_type or content_type.lower() in ("application/octet-stream", "binary/octet-stream"):
        return os.path.splitext(filename or "")[1].lower() in EXT_ALLOWLIST
    return content_type.lower() in ALLOWED_TYPES


# Optional libraries — each is detected at import so missing deps don't crash
# the server (extractors degrade to a friendly marker string instead).
try:
    from docx import Document as _DocxDocument
    HAS_DOCX = True
except Exception:
    HAS_DOCX = False
try:
    import openpyxl
    HAS_XLSX = True
except Exception:
    HAS_XLSX = False
try:
    from pptx import Presentation as _PptxPresentation
    HAS_PPTX = True
except Exception:
    HAS_PPTX = False
try:
    from mutagen import File as _MutagenFile
    HAS_MUTAGEN = True
except Exception:
    HAS_MUTAGEN = False
try:
    from imageio_ffmpeg import get_ffmpeg_exe as _get_ffmpeg_exe
    HAS_FFMPEG = True
except Exception:
    HAS_FFMPEG = False


class FileInfoResponse(BaseModel):
    id: int
    filename: str
    file_size: int
    mime_type: str
    version: int
    created_at: datetime
    chat_id: int | None = None

    model_config = {"from_attributes": True}


@router.post("/upload")
async def upload_file(
    request: Request,
    file: UploadFile = File(...),
    chat_id: int = None,
    db: AsyncSession = Depends(get_db),
):
    user = await get_current_user_from_cookie(request, db)
    await check_limit(user, db, "file_upload")

    content = await file.read()
    if len(content) > MAX_SIZE:
        raise HTTPException(status_code=400, detail="File too large (max 20MB)")

    if file.content_type and not _file_allowed(file.content_type, file.filename):
        raise HTTPException(status_code=400, detail=f"File type not allowed: {file.content_type}")

    ext = os.path.splitext(file.filename or "file")[1]
    stored_name = f"{user.id}_{uuid.uuid4().hex[:12]}{ext}"
    file_path = os.path.join(UPLOAD_DIR, stored_name)

    with open(file_path, "wb") as f:
        f.write(content)

    db_file = UploadedFile(
        user_id=user.id,
        chat_id=chat_id,
        filename=file.filename or "unnamed",
        stored_name=stored_name,
        file_size=len(content),
        mime_type=file.content_type or "application/octet-stream",
    )
    db.add(db_file)
    await db.commit()
    await db.refresh(db_file)

    return {"file": FileInfoResponse.model_validate(db_file)}


@router.get("")
async def list_files(request: Request, chat_id: int = None, db: AsyncSession = Depends(get_db)):
    user = await get_current_user_from_cookie(request, db)
    query = select(UploadedFile).where(UploadedFile.user_id == user.id)
    if chat_id:
        query = query.where(UploadedFile.chat_id == chat_id)
    result = await db.execute(query.order_by(UploadedFile.created_at.desc()).limit(50))
    files = result.scalars().all()
    return {"files": [FileInfoResponse.model_validate(f) for f in files]}


@router.get("/{file_id}/download")
async def download_file(file_id: int, request: Request, db: AsyncSession = Depends(get_db)):
    user = await get_current_user_from_cookie(request, db)
    result = await db.execute(
        select(UploadedFile).where(UploadedFile.id == file_id, UploadedFile.user_id == user.id)
    )
    db_file = result.scalar_one_or_none()
    if not db_file:
        raise HTTPException(status_code=404, detail="File not found")

    file_path = os.path.join(UPLOAD_DIR, db_file.stored_name)
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="File missing from disk")

    return FastAPIFileResponse(
        path=file_path,
        filename=db_file.filename,
        media_type=db_file.mime_type,
    )


@router.get("/{file_id}/read")
async def read_file(file_id: int, request: Request, db: AsyncSession = Depends(get_db)):
    user = await get_current_user_from_cookie(request, db)
    result = await db.execute(
        select(UploadedFile).where(UploadedFile.id == file_id, UploadedFile.user_id == user.id)
    )
    db_file = result.scalar_one_or_none()
    if not db_file:
        raise HTTPException(status_code=404, detail="File not found")

    file_path = os.path.join(UPLOAD_DIR, db_file.stored_name)
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="File missing from disk")

    mime = (db_file.mime_type or "").lower()
    ext = os.path.splitext(db_file.filename or "")[1].lower()

    if mime == "application/pdf" or ext == ".pdf":
        return {"filename": db_file.filename, "type": "text", "source": "pdf", "content": extract_pdf_text(file_path)}

    if mime.startswith("image/"):
        return {
            "filename": db_file.filename,
            "type": "image",
            "source": "image",
            "url": f"/api/files/{file_id}/download",
            "data_url": file_to_data_url(file_path, db_file.mime_type),
        }

    if mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or ext == ".docx":
        return {"filename": db_file.filename, "type": "text", "source": "docx", "content": extract_docx_text(file_path)}

    if mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" or ext in (".xlsx", ".xlsm"):
        return {"filename": db_file.filename, "type": "text", "source": "xlsx", "content": extract_xlsx_text(file_path)}

    if mime == "application/vnd.openxmlformats-officedocument.presentationml.presentation" or ext == ".pptx":
        return {"filename": db_file.filename, "type": "text", "source": "pptx", "content": extract_pptx_text(file_path)}

    if mime.startswith("audio/") or ext in (".mp3", ".wav", ".m4a", ".aac", ".ogg", ".opus", ".flac", ".wma"):
        return {"filename": db_file.filename, "type": "text", "source": "audio", "content": extract_audio_info(file_path)}

    if mime.startswith("video/") or ext in (".mp4", ".webm", ".mov", ".avi", ".mkv", ".3gp", ".3g2"):
        content, frames = extract_video_info(file_path)
        return {
            "filename": db_file.filename,
            "type": "text",
            "source": "video",
            "content": content,
            "frames": frames[:4],
        }

    try:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
        return {"filename": db_file.filename, "type": "text", "source": "text", "content": _safe_text(content)}
    except Exception:
        raise HTTPException(status_code=400, detail="Cannot read this file type")


@router.delete("/{file_id}")
async def delete_file(file_id: int, request: Request, db: AsyncSession = Depends(get_db)):
    user = await get_current_user_from_cookie(request, db)
    result = await db.execute(
        select(UploadedFile).where(UploadedFile.id == file_id, UploadedFile.user_id == user.id)
    )
    db_file = result.scalar_one_or_none()
    if not db_file:
        raise HTTPException(status_code=404, detail="File not found")

    file_path = os.path.join(UPLOAD_DIR, db_file.stored_name)
    if os.path.exists(file_path):
        os.remove(file_path)

    await db.delete(db_file)
    await db.commit()
    return {"message": "Deleted"}


# ---------------------------------------------------------------------------
# Extraction helpers — designed to never raise, always return a string the AI
# can read (or a data-URL for binary content the frontend can pass to vision).
# ---------------------------------------------------------------------------

def _safe_text(text, cap=MAX_TEXT) -> str:
    text = (text or "").strip()
    if len(text) > cap:
        text = text[:cap] + "\n...[truncated]"
    return text


def file_to_data_url(file_path: str, mime: str) -> str:
    try:
        with open(file_path, "rb") as f:
            raw = f.read()
        if len(raw) > 3 * 1024 * 1024:  # keep payloads sane
            return ""
        return f"data:{mime or 'application/octet-stream'};base64,{base64.b64encode(raw).decode()}"
    except Exception:
        return ""


def extract_pdf_text(file_path: str) -> str:
    try:
        result = subprocess.run(
            ["python", "-c", f"""
import sys
sys.path.insert(0, '.')
try:
    import PyPDF2
    with open(r"{file_path}", "rb") as f:
        reader = PyPDF2.PdfReader(f)
        text = ""
        for page in reader.pages:
            text += page.extract_text() or ""
        print(text[:{MAX_TEXT}])
except ImportError:
    print("[PDF extraction requires PyPDF2. Install with: pip install PyPDF2]")
"""],
            capture_output=True, text=True, timeout=30
        )
        return result.stdout or "[Could not extract PDF text]"
    except Exception as e:
        return f"[PDF extraction error: {e}]"


def extract_docx_text(file_path: str) -> str:
    if not HAS_DOCX:
        return "[docx reading not available on this server]"
    try:
        doc = _DocxDocument(file_path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                parts.append(" | ".join(cells))
        return _safe_text("\n".join(parts)) or "[Empty document]"
    except Exception as e:
        return f"[Could not read docx: {e}]"


def extract_xlsx_text(file_path: str) -> str:
    if not HAS_XLSX:
        return "[xlsx reading not available on this server]"
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        parts = []
        try:
            for ws in wb.worksheets:
                parts.append(f"### Sheet: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = ["" if v is None else str(v) for v in row]
                    if any(c.strip() for c in cells):
                        parts.append(" | ".join(cells))
        finally:
            wb.close()
        out = "\n".join(parts)
        return _safe_text(out) or "[Empty spreadsheet]"
    except Exception as e:
        return f"[Could not read xlsx: {e}]"


def extract_pptx_text(file_path: str) -> str:
    if not HAS_PPTX:
        return "[pptx reading not available on this server]"
    try:
        prs = _PptxPresentation(file_path)
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            parts.append(f"--- Slide {i} ---")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text)
        return _safe_text("\n".join(parts)) or "[Empty presentation]"
    except Exception as e:
        return f"[Could not read pptx: {e}]"


def extract_audio_info(file_path: str) -> str:
    raw = {}
    if HAS_MUTAGEN:
        try:
            mf = _MutagenFile(file_path)
            if mf is not None:
                info = getattr(mf, "info", None)
                if info is not None:
                    for attr in ("length", "bitrate", "sample_rate", "channels", "bits_per_sample"):
                        try:
                            v = getattr(info, attr)
                            if v:
                                raw[attr] = v
                        except Exception:
                            pass
                tags = getattr(mf, "tags", None)
                if tags:
                    for key in ("title", "artist", "album", "genre"):
                        try:
                            v = str(tags.get(key, ""))
                            if v:
                                raw[key] = v
                        except Exception:
                            pass
        except Exception:
            pass
    raw.setdefault("format", os.path.splitext(file_path)[1].lstrip(".") or "audio")
    if "length" in raw:
        secs = float(raw.pop("length"))
        raw["duration"] = f"{int(secs // 60)}m {int(secs % 60)}s"
    if "bitrate" in raw:
        try:
            raw["bitrate"] = f"{int(raw['bitrate']) // 1000} kbps"
        except Exception:
            pass

    header = ", ".join(f"{k}={v}" for k, v in raw.items())
    parts = [f"[Audio file metadata: {header}]"]
    transcribed = transcribe_audio(file_path)
    if transcribed:
        parts.append("[Transcription]:\n" + transcribed)
    else:
        parts.append("[No transcript available: audio transcription needs a WHISPER_API_KEY or OPENAI_API_KEY env (or WHISPER_BASE_URL/WHISPER_MODEL overrides)]")
    return _safe_text("\n\n".join(parts))


def transcribe_audio(file_path: str) -> str:
    key = os.getenv("WHISPER_API_KEY") or os.getenv("OPENAI_API_KEY")
    if not key:
        return ""
    base = os.getenv("WHISPER_BASE_URL", "https://api.openai.com/v1").rstrip("/")
    model = os.getenv("WHISPER_MODEL", "whisper-1")
    try:
        with open(file_path, "rb") as f:
            files = {"file": (os.path.basename(file_path), f, "application/octet-stream")}
            data = {"model": model}
            r = httpx.post(
                f"{base}/audio/transcriptions",
                headers={"Authorization": f"Bearer {key}"},
                files=files,
                data=data,
                timeout=120,
            )
        if r.status_code == 200:
            return (r.json().get("text") or "").strip()
    except Exception:
        pass
    return ""


def _ffmpeg_exe():
    if HAS_FFMPEG:
        try:
            return _get_ffmpeg_exe()
        except Exception:
            pass
    return None


def _video_duration(file_path: str, exe: str) -> float:
    try:
        r = subprocess.run([exe, "-i", file_path], capture_output=True, text=True, timeout=30)
        m = re.search(r"Duration:\s*(\d+):(\d+):(\d+(?:\.\d+)?)", r.stderr or "")
        if m:
            h, mi, s = int(m.group(1)), int(m.group(2)), float(m.group(3))
            return h * 3600 + mi * 60 + s
    except Exception:
        pass
    return 0.0


def extract_video_info(file_path: str):
    frames = []
    duration = 0.0
    exe = _ffmpeg_exe()
    if exe:
        duration = _video_duration(file_path, exe)
        points = [duration * f for f in (0.1, 0.5, 0.9)] if duration > 0 else [0.0]
        with tempfile.TemporaryDirectory() as td:
            for i, t in enumerate(points[:3]):
                out = os.path.join(td, f"frame{i}.jpg")
                try:
                    cmd = [
                        exe, "-y", "-ss", f"{t:.2f}", "-i", file_path,
                        "-frames:v", "1", "-q:v", "3",
                        "-vf", "scale=640:-2", out,
                    ]
                    subprocess.run(cmd, capture_output=True, timeout=30)
                    if os.path.exists(out) and os.path.getsize(out) > 0:
                        with open(out, "rb") as f:
                            frames.append("data:image/jpeg;base64," + base64.b64encode(f.read()).decode())
                except Exception:
                    continue

    fmt = os.path.splitext(file_path)[1].lstrip(".") or "unknown"
    meta = f"[Video metadata: format={fmt}"
    if duration:
        meta += f", duration={int(duration // 60)}m {int(duration % 60)}s"
    meta += "]"

    parts = [meta]
    transcribed = transcribe_audio(file_path)
    if transcribed:
        parts.append("[Audio track transcription]:\n" + transcribed)
    else:
        parts.append("[No transcript available: audio transcription needs a WHISPER_API_KEY or OPENAI_API_KEY env]")
    if frames:
        parts.append(f"[Extracted {len(frames)} key frame(s) attached for context]")
    return _safe_text("\n\n".join(parts)), frames